import streamlit as st
import requests
from fpdf import FPDF
from io import BytesIO
from docx import Document
import PyPDF2
import re

# =============== CONFIG =================
TOGETHER_API_KEY = "f344222815e6430930b7a8b673cc5f4ca00072e5021d4e1eb0054fd5899bcfdc"

# =============== HELPER FUNCTIONS ================

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return "\n".join(full_text)

def extract_text_from_pdf(pdf_file):
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    text = []
    for page in pdf_reader.pages:
        text.append(page.extract_text())
    return "\n".join(text)

def get_recommendation(user_profile):
    headers = {
        "Authorization": f"Bearer {TOGETHER_API_KEY}",
        "Content-Type": "application/json"
    }

    prompt = f"""
You are a professional career counselor. Based on the following user profile, suggest the top 3 career paths. For each role, also suggest 2 skills or tools to learn, top 3 companies hiring for that role, and the average salary.

User Profile:
{user_profile}

Return the result as:
- Recommended Careers:
1. Role: ...
   Skills to Learn: ..., ...
   Top Companies: ..., ..., ...
   Average Salary: ...

2. Role: ...
   Skills to Learn: ..., ...
   Top Companies: ..., ..., ...
   Average Salary: ...

3. Role: ...
   Skills to Learn: ..., ...
   Top Companies: ..., ..., ...
   Average Salary: ...
"""

    payload = {
        "model": "mistralai/Mixtral-8x7B-Instruct-v0.1",
        "prompt": prompt,
        "max_tokens": 600,
        "temperature": 0.7
    }

    try:
        response = requests.post("https://api.together.xyz/v1/completions", headers=headers, json=payload)
        result = response.json()
        return result['choices'][0]['text'].strip()
    except Exception as e:
        return f"Error: {e}"

def create_pdf_report(user_profile, recommendations, skills_to_learn, skill_gaps):
    pdf = FPDF()
    pdf.add_page()
    # Use built-in font instead of external
    pdf.set_font("Arial", size=12)

    pdf.set_text_color(0, 0, 128)
    pdf.cell(200, 10, txt="Career Recommendation Report", ln=True, align='C')

    pdf.set_text_color(0, 0, 0)
    pdf.ln(10)
    pdf.multi_cell(0, 10, f"User Profile:\n{user_profile}")

    pdf.ln(5)
    pdf.multi_cell(0, 10, f"Recommendations:\n{recommendations}")

    if skills_to_learn:
        pdf.ln(5)
        pdf.multi_cell(0, 10, f"Skills to Learn:\n{', '.join(skills_to_learn)}")

    if skill_gaps:
        pdf.ln(5)
        pdf.multi_cell(0, 10, f"Skill Gap Analysis:\n{skill_gaps}")

    pdf_bytes = pdf.output(dest='S').encode('utf-8')
    return pdf_bytes


# ===================== STREAMLIT UI =======================

st.set_page_config(page_title="🎯 Career Path Recommender", page_icon="🎯", layout="centered")
st.title("🎯 AI Career Path Recommender")
st.markdown("Upload your resume and provide your academic info to get personalized career suggestions.")

with st.form(key='user_info_form'):
    col1, col2 = st.columns(2)
    with col1:
        name = st.text_input("Full Name")
        highest_degree = st.selectbox("Highest Degree Obtained", ["High School", "Associate", "Bachelor's", "Master's", "PhD", "Other"])
        field_of_study = st.text_input("Field of Study / Major")
    with col2:
        university = st.text_input("University / College")
        graduation_year = st.number_input("Year of Graduation", min_value=2020, max_value=2100, step=1)
        gpa = st.text_input("GPA (Optional)")

    resume_file = st.file_uploader("Upload Your Resume (PDF or DOCX)", type=["pdf", "docx"])
    submit_button = st.form_submit_button(label='Get Career Recommendation')

if submit_button:
    if not name or not highest_degree or not field_of_study or not university or not graduation_year:
        st.warning("Please fill all the academic information fields.")
    elif not resume_file:
        st.warning("Please upload your resume file.")
    else:
        if resume_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            resume_text = extract_text_from_docx(resume_file)
        elif resume_file.type == "application/pdf":
            resume_text = extract_text_from_pdf(resume_file)
        else:
            st.error("Unsupported file type.")
            resume_text = ""

        full_profile = (
            f"Name: {name}\n"
            f"Highest Degree: {highest_degree}\n"
            f"Field of Study: {field_of_study}\n"
            f"University: {university}\n"
            f"Graduation Year: {int(graduation_year)}\n"
            f"GPA: {gpa}\n\n"
            f"Resume Text:\n{resume_text}"
        )

        with st.spinner("Analyzing your profile with AI..."):
            recommendation = get_recommendation(full_profile)

        st.subheader("Career Recommendations")
        st.text_area("AI's Suggestion:", recommendation, height=300)

        role_blocks = re.findall(r"\d+\.\s*Role: (.*?)\n\s*Skills to Learn: (.*?)\n\s*Top Companies: (.*?)\n\s*Average Salary: (.*?)\n", recommendation, re.DOTALL)

        if role_blocks:
            role_names = [r[0] for r in role_blocks]
            selected_role = st.selectbox("Select a Role for Company & Salary Info", role_names)

            for role, skills, companies, salary in role_blocks:
                if role == selected_role:
                    st.markdown(f"### Companies hiring for **{role}**:")
                    for company in [c.strip() for c in companies.split(',')]:
                        st.write(f"- {company}")
                    st.markdown(f"**Average Salary**: {salary.strip()}")
                    break

        pdf_bytes = create_pdf_report(full_profile, recommendation, [], "")
        pdf_file = BytesIO(pdf_bytes)

        st.download_button(
            label="📥 Download Recommendation Report (PDF)",
            data=pdf_file,
            file_name="career_recommendation_report.pdf",
            mime="application/pdf"
        )

        st.markdown("---")
        st.caption("Developed by Rachana")
